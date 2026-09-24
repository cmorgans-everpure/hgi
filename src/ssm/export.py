"""Write the model to a formula-driven Excel workbook and JSON."""
from __future__ import annotations
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BLUE = Font(color="0000FF")                       # hard-coded inputs
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="1F2A44"); HDRF = Font(bold=True, color="FFFFFF")
USD = '$#,##0;($#,##0);"-"'; NUM = '#,##0;(#,##0);"-"'; PCT = '0.0%'; DEC = '0.00'


def _hdr(ws, row, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v); c.fill = HDR; c.font = HDRF
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x


def to_excel(model: dict, recon: dict, evidence: list[dict], path: Path) -> Path:
    wb = Workbook()
    inp = model["inputs"]

    # ---------- Assumptions ----------
    a = wb.active; a.title = "Assumptions"
    _hdr(a, 1, ["Input", "Value", "Provenance"])
    sp_, net, fe = model["space"], model["networking"], model["finance_esg"]
    pv = model["provenance"]
    rows = [
        ("PUE", model["power_cooling"]["pue"], DEC, pv.get("pue")),                                   # 2
        ("Electricity $/kWh", model["power_cooling"]["usd_per_kwh"], '$0.000', pv.get("usd_per_kwh")), # 3
        ("Refresh cycle (years)", inp["refresh_years"], NUM, "assumption"),                          # 4
        ("Annual support % of acquisition", inp["support_pct"], PCT, "assumption"),                  # 5
        ("Cloud $/GB-month", inp["cloud_usd_per_gb_month"], '$0.0000', "assumption"),                # 6
        ("Usable TB per storage admin FTE", inp["tb_per_admin_fte"], NUM, "assumption"),             # 7
        ("Loaded cost per admin FTE", inp["admin_fte_loaded_cost"], USD, "assumption"),              # 8
        ("Revenue (USD)", inp["revenue_usd"], USD, pv.get("revenue")),                               # 9
        ("IT spend % of revenue", inp["it_pct_revenue"], PCT, f"industry: {model['industry_used']}"), # 10
        ("IT spend (USD)", None, USD, "formula unless disclosed"),                                   # 11
        ("Storage % of IT - low", inp["storage_pct_it"]["low"], PCT, "assumption"),                  # 12
        ("Storage % of IT - mid", inp["storage_pct_it"]["mid"], PCT, "assumption"),                  # 13
        ("Storage % of IT - high", inp["storage_pct_it"]["high"], PCT, "assumption"),                # 14
        ("Disclosed cloud storage spend", inp["cloud_spend_disclosed"], USD, "evidence (blank = modelled)"),  # 15
        ("Space cost $/RU/month", sp_["usd_per_ru_month"], USD, pv.get("usd_per_ru_month")),         # 16
        ("Rack overhead RU per 42U rack", model["inputs"]["rack_overhead_ru"], NUM, "assumption"),    # 17
        ("Storage network switches (count)", net["switches"], NUM, f"{net['sites']} site(s)"),       # 18
        ("Avg switch price", net["capex"] / net["switches"] if net["switches"] else 0, USD, "assumption"),  # 19
        ("Switch support % / yr", model["inputs"]["switch_support_pct"], PCT, "assumption"),          # 20
        ("Annual data growth %", model["baseline"]["growth_pct"], PCT, pv.get("data_growth_pct")),    # 21
        ("WACC / discount rate", fe["wacc_pct"], PCT, pv.get("wacc")),                                # 22
        ("Grid kgCO2e per kWh", fe["grid_kgco2e_per_kwh"], '0.000', pv.get("grid_factor")),           # 23
        ("Rack units override (blank = computed)", sp_["rack_units"] if sp_["basis"].startswith("evidence") else None, NUM, sp_["basis"]),  # 24
    ]
    for i, (k, v, fmt, prov) in enumerate(rows, 2):
        a.cell(row=i, column=1, value=k)
        c = a.cell(row=i, column=2, value=v); c.number_format = fmt; c.font = BLUE
        a.cell(row=i, column=3, value=prov)
    it_from_ev = inp.get("it_spend_disclosed")
    a["B11"] = it_from_ev if it_from_ev is not None else '=IF(B9="","",B9*B10)'
    a["B11"].font = BLUE if it_from_ev is not None else Font(color="000000")
    _widths(a, [38, 18, 60])

    # ---------- Spend Model ----------
    m = wb.create_sheet("Spend Model")
    _hdr(m, 1, ["Tier", "Usable TB", "$/TB (acq.)", "W/TB", "Acquisition value", "Hardware (annualized)",
                "Maint. & support", "IT kWh / yr", "Power $", "Cooling & facility $", "Cloud $", "Annual total",
                "Price basis", "TB per RU", "Rack units"])
    tiers = ["performance", "general", "capacity", "archive"]
    for r, t in enumerate(tiers, 2):
        L = model["tier_lines"][t]
        m.cell(row=r, column=1, value=t)
        for col, val, fmt in ((2, L["tb"], NUM), (3, L["usd_per_tb"], USD), (4, L["watts_per_tb"], DEC), (14, L["tb_per_ru"], NUM)):
            c = m.cell(row=r, column=col, value=round(val, 4)); c.number_format = fmt; c.font = BLUE
        f = {5: f"=B{r}*C{r}", 6: f"=E{r}/Assumptions!$B$4", 7: f"=E{r}*Assumptions!$B$5", 8: f"=B{r}*D{r}*8.76",
             9: f"=H{r}*Assumptions!$B$3", 10: f"=I{r}*(Assumptions!$B$2-1)", 11: 0, 12: f"=F{r}+G{r}+I{r}+J{r}+K{r}",
             15: f"=IF(N{r}=0,0,B{r}/N{r}*42/(42-Assumptions!$B$17))"}
        for col, fx in f.items():
            c = m.cell(row=r, column=col, value=fx); c.number_format = NUM if col in (8, 15) else USD
        m.cell(row=r, column=13, value=model["pricing"]["basis"][t])
    m["A6"] = "cloud"; m["B6"] = round(model["capacity"]["by_tier"]["cloud"], 4); m["B6"].font = BLUE
    m["K6"] = '=IF(Assumptions!B15<>"",Assumptions!B15,B6*1024*Assumptions!B6*12)'; m["L6"] = "=K6"
    m["A7"] = "admin labor"; m["L7"] = "=SUM(B2:B5)/Assumptions!B7*Assumptions!B8"
    m["A8"] = "data centre space"; m["O8"] = '=IF(Assumptions!B24<>"",Assumptions!B24,SUM(O2:O5))'
    m["L8"] = "=O8*Assumptions!B16*12"
    m["A9"] = "networking"; m["L9"] = "=Assumptions!B18*Assumptions!B19/Assumptions!B4+Assumptions!B18*Assumptions!B19*Assumptions!B20"
    m["A10"] = "TOTAL (bottom-up)"; m["A10"].font = BOLD
    for col in "BEFGHIJKL":
        m[f"{col}10"] = f"=SUM({col}2:{col}9)"; m[f"{col}10"].font = BOLD
    m["O10"] = "=O8"; m["O10"].font = BOLD
    for rr in range(6, 11):
        for col in "EFGIJKL": m[f"{col}{rr}"].number_format = USD
        for col in "BHO": m[f"{col}{rr}"].number_format = NUM
    m["A12"] = "Capacity basis"; m["B12"] = model["capacity"]["basis"]
    for i, (k, v) in enumerate(model["capacity"]["methods"].items(), 13):
        m.cell(row=i, column=1, value=f"  capacity via {k}"); c = m.cell(row=i, column=2, value=v); c.number_format = NUM
    _widths(m, [20, 14, 12, 8, 16, 16, 16, 14, 14, 16, 14, 16, 50, 10, 11])

    # ---------- Current-State Baseline (5-yr, value-prop format) ----------
    b = wb.create_sheet("Current-State Baseline")
    yrs = model["baseline"]["years"]
    b["A1"] = f"Current-state {yrs}-year cost of ownership - incumbent environment"; b["A1"].font = Font(bold=True, size=13)
    b["A2"] = "Capacity-driven lines grow at the annual data-growth rate (Assumptions!B21). Networking is flat."
    _hdr(b, 4, ["Category"] + [f"Year {y}" for y in range(1, yrs + 1)] + [f"{yrs}-yr total"])
    SM = "'Spend Model'!"
    g = "(1+Assumptions!$B$21)^"
    lines_ = [("Platform (HW, SW & support)", f"({SM}$F$10+{SM}$G$10)", True),
              ("Networking", f"{SM}$L$9", False),
              ("Power & Cooling", f"({SM}$I$10+{SM}$J$10)", True),
              ("Data Centre Space", f"{SM}$L$8", True),
              ("Admin labor", f"{SM}$L$7", True),
              ("Cloud storage", f"{SM}$K$10", True)]
    last = get_column_letter(yrs + 1); totc = get_column_letter(yrs + 2)
    for i, (lab, base, grows) in enumerate(lines_, 5):
        b.cell(row=i, column=1, value=lab)
        for y in range(yrs):
            c = b.cell(row=i, column=y + 2, value=f"={base}*{g}{y}" if grows else f"={base}"); c.number_format = USD
        c = b.cell(row=i, column=yrs + 2, value=f"=SUM(B{i}:{last}{i})"); c.number_format = USD; c.font = BOLD
    tr = 5 + len(lines_)
    b.cell(row=tr, column=1, value="Total current state").font = BOLD
    for col in range(2, yrs + 3):
        L_ = get_column_letter(col); c = b.cell(row=tr, column=col, value=f"=SUM({L_}5:{L_}{tr-1})"); c.number_format = USD; c.font = BOLD
    r0 = tr + 2
    b.cell(row=r0, column=1, value="Storage facility energy (kWh/yr, incl. PUE)"); c = b.cell(row=r0, column=2, value=f"={SM}H10*Assumptions!B2"); c.number_format = NUM
    b.cell(row=r0 + 1, column=1, value="Storage emissions (tCO2e/yr)"); c = b.cell(row=r0 + 1, column=2, value=f"=B{r0}*Assumptions!B23/1000"); c.number_format = NUM
    b.cell(row=r0 + 2, column=1, value="Rack units (storage)"); c = b.cell(row=r0 + 2, column=2, value=f"={SM}O8"); c.number_format = NUM
    b.cell(row=r0 + 3, column=1, value="WACC for NPV"); c = b.cell(row=r0 + 3, column=2, value="=Assumptions!B22"); c.number_format = PCT
    b.cell(row=r0 + 4, column=1, value=f"NPV of {yrs}-yr current-state cost"); c = b.cell(row=r0 + 4, column=2, value=f"=NPV(Assumptions!B22,B{tr}:{last}{tr})"); c.number_format = USD
    _widths(b, [40] + [15] * (yrs + 1))

    # ---------- Capacity Profile ----------
    cp = wb.create_sheet("Capacity Profile")
    prof = model["capacity_profile"]
    _hdr(cp, 1, ["Protocol", "Usable TB", "Usable %", "Utilisation %", "Data-reduction ratio", "Raw TB", "Stored data (effective TB)"])
    for i, (p, row) in enumerate(prof["by_protocol"].items(), 2):
        cp.cell(row=i, column=1, value=p)
        for col, val, fmt in ((2, row["usable_tb"], NUM), (3, prof["usable_pct"], PCT), (4, prof["utilization_pct"], PCT), (5, row["drr"], DEC)):
            c = cp.cell(row=i, column=col, value=val); c.number_format = fmt; c.font = BLUE
        c = cp.cell(row=i, column=6, value=f"=IF(C{i}=0,0,B{i}/C{i})"); c.number_format = NUM
        c = cp.cell(row=i, column=7, value=f"=B{i}*D{i}*E{i}"); c.number_format = NUM
    n = len(prof["by_protocol"]) + 2
    cp.cell(row=n, column=1, value="Total").font = BOLD
    for col in "BFG":
        c = cp[f"{col}{n}"]; c.value = f"=SUM({col}2:{col}{n-1})"; c.number_format = NUM; c.font = BOLD
    _widths(cp, [14, 14, 11, 13, 20, 14, 24])

    # ---------- Value-Prop Readiness ----------
    rd = wb.create_sheet("Value-Prop Readiness")
    _hdr(rd, 1, ["Input", "Critical", "Status", "Detail / value used", "Evidence IDs"])
    for i, r in enumerate(model["readiness"], 2):
        for j, v in enumerate([r["input"], "yes" if r["critical"] else "", r["status"], r["detail"], ", ".join(r["evidence"])], 1):
            rd.cell(row=i, column=j, value=v)
    _widths(rd, [36, 9, 18, 70, 30])

    # ---------- Refresh & Narrative ----------
    rn = wb.create_sheet("Refresh & Narrative")
    _hdr(rn, 1, ["System / item", "Install year", "Age (yrs)", "Refresh due", "Status", "Evidence"])
    for i, r in enumerate(model["refresh"], 2):
        for j, k in enumerate(["system", "install_year", "age_years", "refresh_due", "status", "evidence_id"], 1):
            rn.cell(row=i, column=j, value=r[k])
    r0 = len(model["refresh"]) + 4
    for k, items in model["narrative"].items():
        rn.cell(row=r0, column=1, value=k.replace("_", " ").title()).font = BOLD; r0 += 1
        for it in items or [{"value": "- none found -", "source": "", "id": ""}]:
            rn.cell(row=r0, column=1, value=it["value"]); rn.cell(row=r0, column=5, value=it["id"]); rn.cell(row=r0, column=6, value=it["source"]); r0 += 1
        r0 += 1
    _widths(rn, [60, 12, 10, 12, 16, 50])

    # ---------- Summary ----------
    s = wb.create_sheet("Summary", 0)
    co = model["company"]; est = model["estimates"]
    s["A1"] = f"Storage spend model - {co['name']}"; s["A1"].font = Font(bold=True, size=14)
    s["A2"] = f"Industry: {model['industry_used']}   Region: {co.get('region')}   Ticker: {co.get('ticker') or '-'}"
    _hdr(s, 4, ["Method", "Annual spend", "Weight"])
    meth = est["methods_used"]
    s["A5"] = "Bottom-up"; s["B5"] = "='Spend Model'!L10"; s["C5"] = meth.get("bottom_up", {}).get("weight", 0)
    s["A6"] = "Top-down (mid)"; s["B6"] = '=IF(Assumptions!B11="",0,Assumptions!B11*Assumptions!B13)'; s["C6"] = meth.get("top_down", {}).get("weight", 0)
    s["A7"] = "Disclosed storage spend"; s["B7"] = est["disclosed_spend"] or 0; s["C7"] = meth.get("disclosed_spend", {}).get("weight", 0)
    s["A8"] = "TRIANGULATED ESTIMATE"; s["A8"].font = BOLD
    s["B8"] = '=IF(SUM(C5:C7)=0,"",SUMPRODUCT(B5:B7,C5:C7)/SUM(C5:C7))'; s["B8"].font = BOLD
    s["A9"] = "Range low"; s["B9"] = est["range"][0]; s["A10"] = "Range high"; s["B10"] = est["range"][1]
    s["A11"] = "Confidence (0-100)"; s["B11"] = est["confidence"]
    for rr in range(5, 11): s[f"B{rr}"].number_format = USD
    for rr in (5, 6, 7): s[f"C{rr}"].font = BLUE; s[f"C{rr}"].number_format = DEC
    _hdr(s, 13, ["Cost category", "Annual $", "Share"])
    cats = [("Hardware (annualized)", "='Spend Model'!F10"), ("Maintenance & support", "='Spend Model'!G10"),
            ("Power", "='Spend Model'!I10"), ("Cooling & facility", "='Spend Model'!J10"),
            ("Cloud storage", "='Spend Model'!K10"), ("Admin labor", "='Spend Model'!L7"),
            ("Data centre space", "='Spend Model'!L8"), ("Networking", "='Spend Model'!L9")]
    tot = 14 + len(cats)
    for i, (lab, fx) in enumerate(cats, 14):
        s[f"A{i}"] = lab; s[f"B{i}"] = fx; s[f"C{i}"] = f"=B{i}/$B${tot}"
    s[f"A{tot}"] = "Bottom-up total"; s[f"B{tot}"] = f"=SUM(B14:B{tot-1})"; s[f"A{tot}"].font = BOLD
    for rr in range(14, tot + 1): s[f"B{rr}"].number_format = USD; s[f"C{rr}"].number_format = PCT
    r = tot + 2
    crit_missing = [x["input"] for x in model["readiness"] if x["critical"] and x["status"] in ("missing", "assumed")]
    for lab, val in (("Primary storage vendor", model["environment"]["primary_vendor"]),
                     ("Vendors identified", ", ".join(model["environment"]["vendors_ranked"])),
                     (f"{yrs}-yr current-state cost", "='Current-State Baseline'!" + f"{totc}{tr}"),
                     ("Reconciliation", recon["summary"]),
                     ("Critical inputs to confirm", "; ".join(crit_missing) or "none")):
        s[f"A{r}"] = lab; s[f"B{r}"] = val; r += 1
    s[f"B{r-3}"].number_format = USD
    s[f"A{r+1}"] = "Blue = hard-coded input; black = formula. All inputs traced on the Evidence and Assumptions sheets."
    _widths(s, [34, 20, 10])

    # ---------- Environment ----------
    e = wb.create_sheet("Environment")
    _hdr(e, 1, ["Attribute", "Value", "Storage tier", "Trust tier", "Verified", "Evidence ID", "Source"])
    for i, it in enumerate(model["environment"]["items"], 2):
        for j, k in enumerate(["field", "value", "tier", "trust_tier", "verified", "evidence_id", "source"], 1):
            e.cell(row=i, column=j, value=str(it[k]) if it[k] is not None else "")
    _widths(e, [24, 40, 14, 10, 10, 12, 70])

    # ---------- Reconciliation ----------
    rc = wb.create_sheet("Reconciliation")
    _hdr(rc, 1, ["Check", "Ours", "Reference", "Source", "Variance", "Status", "Note"])
    for i, c in enumerate(recon["checks"], 2):
        for j, k in enumerate(["check", "ours", "reference", "source", "variance_pct", "status", "note"], 1):
            cell = rc.cell(row=i, column=j, value=c[k])
            if k in ("ours", "reference") and isinstance(c[k], (int, float)): cell.number_format = USD
            if k == "variance_pct": cell.number_format = PCT
    r0 = len(recon["checks"]) + 3
    rc.cell(row=r0, column=1, value="Vendors confirmed by HGI").font = BOLD; rc.cell(row=r0, column=2, value=", ".join(recon["vendors"]["confirmed_by_hgi"]))
    rc.cell(row=r0 + 1, column=1, value="Found only by research").font = BOLD; rc.cell(row=r0 + 1, column=2, value=", ".join(recon["vendors"]["found_only_by_research"]))
    rc.cell(row=r0 + 2, column=1, value="Only in HGI").font = BOLD; rc.cell(row=r0 + 2, column=2, value=", ".join(recon["vendors"]["only_in_hgi"]))
    _widths(rc, [30, 18, 18, 18, 10, 14, 60])

    # ---------- Evidence ----------
    ev = wb.create_sheet("Evidence")
    cols = ["id", "field", "value", "unit", "tier_hint", "trust_tier", "verified", "source_type", "source_title",
            "published_date", "source_url", "quote", "verification_note", "found_by"]
    _hdr(ev, 1, cols)
    for i, x in enumerate(evidence, 2):
        for j, k in enumerate(cols, 1):
            v = x.get(k); ev.cell(row=i, column=j, value=str(v) if v is not None else "")
    _widths(ev, [8, 22, 28, 8, 12, 8, 8, 16, 30, 12, 50, 80, 30, 18])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def to_json(model, recon, evidence, path: Path) -> Path:
    path.write_text(json.dumps({"model": model, "reconciliation": recon, "evidence": evidence}, indent=2, default=str))
    return path
